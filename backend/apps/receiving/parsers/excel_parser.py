"""
Парсер Excel-файлов приёмки от лабораторий.

Формат (решение 1.49 + уточнения receiving-1a / receiving-1a-2):
- Один файл = одна лаборатория.
- Один лист = один участок.
- Строки = Н/З.
- Столбцы = типы исследования, разбитые на группы.
- Каждая группа типов заканчивается колонкой «Дата».
- Ячейки = количество тары; могут быть составными («1+1»).
- Даты в ячейке «Дата» могут быть списком через запятую
  («13.08,02.09,25.05»). Без года — подставляется текущий.

Файл может содержать «шапку» из одной-двух строк до строки
заголовков. Парсер сам находит строку заголовков по маркеру
«№ Наряд-заказа» / «Н/З» / «Наряд».

Сразу после строки заголовков идёт строка **типа тары**
(«Ящики», «Биг-бэг», …). Признак строки типа тары:
ячейка под колонкой Н/З **пустая**, но есть другие значения.
Если под Н/З есть значение — это уже данные, type_row нет.

Парсер не обращается к БД. Список известных кодов типов
передаётся снаружи (`research_type_codes`).
"""

import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Union

from openpyxl import load_workbook

from .types import ParsedCell, ParsedRow, ParsedSheet, ParsedWorkbook

WORK_ORDER_HEADERS = {
    "н/з", "наряд", "наряд-заказ", "наряд заказ", "нз",
    "№ наряд-заказа", "№ наряд заказа",
}
DATE_HEADERS = {"дата", "date"}
NON_TYPE_HEADERS = WORK_ORDER_HEADERS | DATE_HEADERS | {"№", "no", "n"}

# Коды ошибок (структурированные — чтобы UI мог различать).
ERR_MISSING_WORK_ORDER = "MISSING_WORK_ORDER_COLUMN"
ERR_UNKNOWN_TYPE = "UNKNOWN_RESEARCH_TYPE"
ERR_BAD_NUMBER = "BAD_COUNT_VALUE"
ERR_BAD_DATE = "BAD_DATE_VALUE"
ERR_MISSING_CONTAINER_TYPE = "MISSING_CONTAINER_TYPE"

# Максимум строк «шапки» до заголовков (защита от зацикливания).
MAX_HEADER_SEARCH = 20


def _normalize(value) -> str:
    """Нормализация строки: strip, lower, схлопывание пробелов."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip().lower())


def _parse_one_date(s: str, default_year: int) -> date | None:
    """Одна дата. Форматы: YYYY-MM-DD, DD.MM.YYYY, DD.MM.YY,
    DD.MM, DD/MM/YYYY, DD-MM-YYYY."""
    s = s.strip()
    if not s:
        return None

    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except ValueError:
        pass

    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue

    for fmt in ("%d.%m.%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue

    # Без года — парсим вручную (обход DeprecationWarning Python 3.15).
    parts = re.split(r"[./-]", s)
    if len(parts) == 2:
        try:
            day = int(parts[0])
            month = int(parts[1])
            return date(default_year, month, day)
        except (ValueError, TypeError):
            return None

    return None


def _parse_date_list(value, default_year: int | None = None) -> list[date]:
    """Список дат из ячейки (одна / через запятую / через ;)."""
    if default_year is None:
        default_year = date.today().year

    if value is None:
        return []
    if isinstance(value, datetime):
        return [value.date()]
    if isinstance(value, date):
        return [value]
    if not isinstance(value, str):
        return []

    s = value.strip()
    if not s:
        return []

    parts = [p.strip() for p in re.split(r"[,;]", s) if p.strip()]
    result: list[date] = []
    for part in parts:
        d = _parse_one_date(part, default_year)
        if d is not None:
            result.append(d)
    return result


def _parse_count(value) -> int | None:
    """Количество тары (int / float / строка / «1+1»)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        n = int(value)
        return n if n > 0 else None
    if not isinstance(value, str):
        return None

    s = value.strip()
    if not s:
        return None

    parts = [p.strip() for p in s.split("+")]
    total = 0
    for part in parts:
        try:
            total += int(float(part))
        except (ValueError, TypeError):
            return None
    return total if total > 0 else None


def parse_excel(
    source: Union[BytesIO, Path, str],
    research_type_codes: set[str] | None = None,
) -> ParsedWorkbook:
    """Парсит Excel-файл приёмки."""
    wb = load_workbook(source, data_only=True, read_only=True)
    result = ParsedWorkbook()

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parsed_sheet = _parse_sheet(
                ws, sheet_name, research_type_codes, result.errors,
            )
            if parsed_sheet.rows:
                result.sheets.append(parsed_sheet)
    finally:
        wb.close()

    return result


def _parse_sheet(
    ws,
    sheet_name: str,
    research_type_codes: set[str] | None,
    errors: list[dict],
) -> ParsedSheet:
    """Парсит один лист книги."""
    sheet = ParsedSheet(sheet_name=sheet_name, site_name=sheet_name)

    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return sheet

    # 1. Ищем строку заголовков.
    header_idx: int | None = None
    for i, row in enumerate(all_rows[:MAX_HEADER_SEARCH]):
        for cell in row:
            if _normalize(cell) in WORK_ORDER_HEADERS:
                header_idx = i
                break
        if header_idx is not None:
            break

    if header_idx is None:
        errors.append({
            "sheet": sheet_name,
            "error": "Не найдена строка заголовков (Н/З)",
            "code": ERR_MISSING_WORK_ORDER,
        })
        return sheet

    headers = all_rows[header_idx]
    sheet.header_row_number = header_idx + 1

    # 2. Ищем колонку Н/З (нужно ДО определения строки типа тары).
    work_order_idx: int | None = None
    for idx, header in enumerate(headers):
        if _normalize(header) in WORK_ORDER_HEADERS:
            work_order_idx = idx
            break

    if work_order_idx is None:
        errors.append({
            "sheet": sheet_name,
            "row": header_idx + 1,
            "error": "Не найдена колонка Н/З",
            "code": ERR_MISSING_WORK_ORDER,
        })
        return sheet

    # 3. Строка типа тары: следующая после заголовков.
    #    Признак: ячейка под Н/З пустая, но есть другие значения.
    #    Если под Н/З есть значение — это уже данные.
    type_row: tuple | None = None
    data_start_idx = header_idx + 1
    if header_idx + 1 < len(all_rows):
        candidate = all_rows[header_idx + 1]
        wo_cell = (
            candidate[work_order_idx]
            if work_order_idx < len(candidate)
            else None
        )
        wo_is_empty = _normalize(wo_cell) == ""
        has_other_values = any(_normalize(c) for c in candidate)
        if wo_is_empty and has_other_values:
            type_row = candidate
            data_start_idx = header_idx + 2

    # 4. Определяем группы «типы + Дата».
    #    Ошибки MISSING_CONTAINER_TYPE собираем в буфер, чтобы добавить
    #    их в общий список только при наличии данных.
    type_groups: list[
        tuple[list[tuple[int, str, str]], int | None]
    ] = []
    current_group: list[tuple[int, str, str]] = []
    missing_ct_errors: list[dict] = []

    for idx, header in enumerate(headers):
        norm = _normalize(header)
        if not norm:
            continue

        if norm in WORK_ORDER_HEADERS:
            continue

        if norm in DATE_HEADERS:
            type_groups.append((current_group, idx))
            current_group = []
            continue

        if norm in NON_TYPE_HEADERS:
            continue

        # Тип исследования.
        if research_type_codes is not None:
            matched = next(
                (c for c in research_type_codes
                 if _normalize(c) == norm),
                None,
            )
            if matched is None:
                errors.append({
                    "sheet": sheet_name,
                    "row": header_idx + 1,
                    "column": idx,
                    "error": f"Неизвестный тип исследования: {header}",
                    "code": ERR_UNKNOWN_TYPE,
                })
                continue
            type_code = matched
        else:
            type_code = str(header).strip()

        # Тип тары из строки type_row.
        container_type_name = ""
        if type_row is not None and idx < len(type_row):
            ct = type_row[idx]
            if ct is not None:
                container_type_name = str(ct).strip()

        if not container_type_name:
            missing_ct_errors.append({
                "sheet": sheet_name,
                "row": header_idx + 2 if type_row else header_idx + 1,
                "column": idx,
                "error": (
                    f"Не указан тип тары для типа исследования "
                    f"«{header}»"
                ),
                "code": ERR_MISSING_CONTAINER_TYPE,
            })
            continue

        current_group.append((idx, type_code, container_type_name))

    # Хвост без колонки «Дата».
    if current_group:
        type_groups.append((current_group, None))

    type_groups = [(cols, di) for cols, di in type_groups if cols]

    # 5. Данные.
    for row_num, row in enumerate(
        all_rows[data_start_idx:], start=data_start_idx + 1,
    ):
        wo_value = (
            row[work_order_idx] if work_order_idx < len(row) else None
        )
        if _normalize(wo_value) == "":
            continue

        parsed_row = ParsedRow(
            work_order_number=str(wo_value).strip(),
        )

        for type_cols, date_col_idx in type_groups:
            group_dates: list[date] = []
            if date_col_idx is not None and date_col_idx < len(row):
                group_dates = _parse_date_list(row[date_col_idx])

            for col_idx, type_code, ct_name in type_cols:
                if col_idx >= len(row):
                    continue
                raw = row[col_idx]
                count = _parse_count(raw)
                if count is None:
                    continue

                parsed_row.cells.append(ParsedCell(
                    research_type_code=type_code,
                    containers_count=count,
                    container_type_name=ct_name,
                    dates=list(group_dates),
                ))

        if parsed_row.cells:
            sheet.rows.append(parsed_row)

    # 6. Добавляем ошибки про типы тары, только если в файле есть
    #    хотя бы одна строка данных (непустой Н/З). Пустой файл
    #    («только заголовки») — не ошибка.
    has_data = any(
        _normalize(row[work_order_idx]) != ""
        for row in all_rows[data_start_idx:]
        if work_order_idx < len(row)
    )
    if has_data:
        errors.extend(missing_ct_errors)

    return sheet