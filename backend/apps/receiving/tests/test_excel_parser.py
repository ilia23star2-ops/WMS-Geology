"""
Тесты парсера Excel-файлов приёмки.

Покрывают:
- простой файл (1 лист, 2 Н/З);
- пустые ячейки — пропускаются;
- несколько листов;
- неизвестный тип исследования;
- разные форматы даты;
- отсутствие колонки Н/З;
- пустые строки — пропускаются;
- нечисловое значение в ячейке;
- альтернативные названия колонки Н/З;
- нулевое количество — пропускается;
- пустой файл;
- ГРУППЫ: несколько блоков «типы + Дата» (receiving-1a);
- список дат через запятую (receiving-1a);
- составное количество «1+1» (receiving-1a);
- без года — текущий год (receiving-1a);
- группа без колонки «Дата» — dates=[].
- ШАПКА: файл с двумя строками перед заголовками (receiving-1a-2);
- ТИП ТАРЫ: строка типа тары после заголовков (receiving-1a-2);
- тип тары отсутствует — ошибка (receiving-1a-2);
- тип тары на конкретную колонку (Биг-бэг, не Ящики) (receiving-1a-2);
- тип тары короче заголовков (receiving-1a-2).
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook

from apps.receiving.parsers.excel_parser import parse_excel


def _make_xlsx(rows: list[list], sheet_name: str = "Тестовый") -> BytesIO:
    """Создаёт XLSX-файл в памяти с одним листом."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ============================================================
# Базовые сценарии
# ============================================================
def test_parse_simple_file():
    """Простой файл: 1 лист, 2 Н/З, типы ШЛ и ХА, общая дата."""
    f = _make_xlsx([
        ["Наряд", "ШЛ", "ХА", "Дата"],
        ["", "Ящики", "Ящики", ""],
        ["Тест001", 5, 3, "2026-10-08"],
        ["Тест002", 2, 4, "2026-10-09"],
    ])
    result = parse_excel(f)
    assert len(result.sheets) == 1
    assert len(result.sheets[0].rows) == 2
    row0 = result.sheets[0].rows[0]
    assert row0.work_order_number == "Тест001"
    assert len(row0.cells) == 2
    assert row0.cells[0].research_type_code == "ШЛ"
    assert row0.cells[0].containers_count == 5
    assert row0.cells[0].container_type_name == "Ящики"
    assert row0.cells[0].dates == [date(2026, 10, 8)]


def test_parse_with_empty_cells():
    """Пустые ячейки — пропускаются."""
    f = _make_xlsx([
        ["Наряд", "ШЛ", "ХА", "Дата"],
        ["", "Ящики", "Ящики", ""],
        ["Тест001", 5, None, "2026-10-08"],
    ])
    result = parse_excel(f)
    assert len(result.sheets[0].rows[0].cells) == 1
    assert result.sheets[0].rows[0].cells[0].research_type_code == "ШЛ"


def test_parse_multiple_sheets():
    """Несколько листов — несколько участков."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Тестовый"
    ws1.append(["Наряд", "ШЛ", "Дата"])
    ws1.append(["", "Ящики", ""])
    ws1.append(["Тест001", 5, "2026-10-08"])

    ws2 = wb.create_sheet("Северный")
    ws2.append(["Наряд", "ШЛ", "Дата"])
    ws2.append(["", "Ящики", ""])
    ws2.append(["Сев001", 3, "2026-10-08"])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    result = parse_excel(buf)
    assert len(result.sheets) == 2
    assert result.sheets[0].sheet_name == "Тестовый"
    assert result.sheets[1].sheet_name == "Северный"


def test_parse_with_unknown_type():
    """Неизвестный тип — записывается в errors, парсинг продолжается."""
    f = _make_xlsx([
        ["Наряд", "ШЛ", "НЕИЗВ", "Дата"],
        ["", "Ящики", "Ящики", ""],
        ["Тест001", 5, 3, "2026-10-08"],
    ])
    result = parse_excel(f, research_type_codes={"ШЛ", "ХА"})
    assert len(result.errors) == 1
    assert "НЕИЗВ" in result.errors[0]["error"]
    assert result.errors[0]["code"] == "UNKNOWN_RESEARCH_TYPE"
    assert len(result.sheets[0].rows[0].cells) == 1


def test_parse_with_different_date_formats():
    """Разные форматы даты обрабатываются."""
    f = _make_xlsx([
        ["Наряд", "ШЛ", "Дата"],
        ["", "Ящики", ""],
        ["Тест001", 5, "08.10.2026"],
        ["Тест002", 3, "2026-10-09"],
    ])
    result = parse_excel(f)
    assert result.sheets[0].rows[0].cells[0].dates == [date(2026, 10, 8)]
    assert result.sheets[0].rows[1].cells[0].dates == [date(2026, 10, 9)]


def test_parse_no_work_order_column():
    """Без колонки Н/З — ошибка, лист пропускается."""
    f = _make_xlsx([
        ["Что-то", "ШЛ", "Дата"],
        ["x", 5, "2026-10-08"],
    ])
    result = parse_excel(f)
    assert len(result.errors) >= 1
    assert "Н/З" in result.errors[0]["error"] or "заголовков" in result.errors[0]["error"]
    assert result.sheets == []


def test_parse_skips_empty_rows():
    """Пустые строки пропускаются."""
    f = _make_xlsx([
        ["Наряд", "ШЛ", "Дата"],
        ["", "Ящики", ""],
        ["Тест001", 5, "2026-10-08"],
        [None, None, None],
        ["Тест002", 3, "2026-10-09"],
    ])
    result = parse_excel(f)
    assert len(result.sheets[0].rows) == 2


def test_parse_total_containers():
    """Свойства total_containers и total_rows считают суммы."""
    f = _make_xlsx([
        ["Наряд", "ШЛ", "ХА", "Дата"],
        ["", "Ящики", "Ящики", ""],
        ["Тест001", 5, 3, "2026-10-08"],
        ["Тест002", 2, 4, "2026-10-09"],
    ])
    result = parse_excel(f)
    assert result.total_containers == 14
    assert result.total_rows == 2


def test_parse_invalid_number_in_cell():
    """Нечисловое значение — ошибка, парсинг продолжается."""
    f = _make_xlsx([
        ["Наряд", "ШЛ", "Дата"],
        ["", "Ящики", ""],
        ["Тест001", "abc", "2026-10-08"],
        ["Тест002", 3, "2026-10-09"],
    ])
    result = parse_excel(f)
    assert len(result.sheets[0].rows) == 1
    assert result.sheets[0].rows[0].work_order_number == "Тест002"


def test_parse_alt_work_order_headers():
    """Альтернативные названия колонки Н/З."""
    for header in ["Н/З", "Наряд-Заказ", "наряд заказ", "НЗ", "№ Наряд-заказа"]:
        f = _make_xlsx([
            [header, "ШЛ", "Дата"],
            ["", "Ящики", ""],
            ["Тест001", 5, "2026-10-08"],
        ])
        result = parse_excel(f)
        assert len(result.sheets[0].rows) == 1, (
            f"Header '{header}' не распознан"
        )


def test_parse_zero_count_skipped():
    """Нулевое количество — ячейка пропускается."""
    f = _make_xlsx([
        ["Наряд", "ШЛ", "ХА", "Дата"],
        ["", "Ящики", "Ящики", ""],
        ["Тест001", 0, 5, "2026-10-08"],
    ])
    result = parse_excel(f)
    assert len(result.sheets[0].rows[0].cells) == 1
    assert result.sheets[0].rows[0].cells[0].research_type_code == "ХА"


def test_parse_empty_file():
    """Файл без данных — пустой результат."""
    f = _make_xlsx([["Наряд", "ШЛ", "Дата"]])
    result = parse_excel(f)
    assert result.sheets == []
    assert result.errors == []


# ============================================================
# Группы «типы + Дата» (receiving-1a)
# ============================================================
def test_parse_two_type_groups():
    """Два блока типов с разными датами."""
    f = _make_xlsx([
        ["Наряд", "ПА", "Дата", "ГД", "Дата"],
        ["", "Ящики", "", "Ящики", ""],
        ["TST00469", 3, "2026-09-08", 4, "2026-09-15"],
    ])
    result = parse_excel(f)
    assert len(result.errors) == 0
    cells = result.sheets[0].rows[0].cells
    assert len(cells) == 2

    pa = next(c for c in cells if c.research_type_code == "ПА")
    gd = next(c for c in cells if c.research_type_code == "ГД")
    assert pa.containers_count == 3
    assert pa.dates == [date(2026, 9, 8)]
    assert gd.containers_count == 4
    assert gd.dates == [date(2026, 9, 15)]


def test_parse_date_list_in_cell():
    """Ячейка «Дата» со списком через запятую → список дат."""
    f = _make_xlsx([
        ["Наряд", "ГД", "Дата"],
        ["", "Ящики", ""],
        ["TST00472", 3, "13.08,02.09,25.05"],
    ])
    result = parse_excel(f)
    cells = result.sheets[0].rows[0].cells
    assert len(cells) == 1
    dates = cells[0].dates
    assert len(dates) == 3
    assert dates[0].month == 8 and dates[0].day == 13
    assert dates[1].month == 9 and dates[1].day == 2
    assert dates[2].month == 5 and dates[2].day == 25


def test_parse_composite_count_1plus1():
    """Составное количество «1+1» = 2."""
    f = _make_xlsx([
        ["Наряд", "ГД", "Дата"],
        ["", "Ящики", ""],
        ["TST00472", "1+1", "2026-09-13"],
    ])
    result = parse_excel(f)
    cells = result.sheets[0].rows[0].cells
    assert len(cells) == 1
    assert cells[0].containers_count == 2


def test_parse_date_without_year_uses_current():
    """Дата без года — подставляется текущий год."""
    current_year = date.today().year
    f = _make_xlsx([
        ["Наряд", "ГД", "Дата"],
        ["", "Ящики", ""],
        ["TST00472", 3, "13.08"],
    ])
    result = parse_excel(f)
    d = result.sheets[0].rows[0].cells[0].dates[0]
    assert d == date(current_year, 8, 13)


def test_parse_group_without_date_column():
    """Группа типов без колонки «Дата» → dates=[]."""
    f = _make_xlsx([
        ["Наряд", "ШЛ", "ХА"],
        ["", "Ящики", "Ящики"],
        ["Тест001", 5, 3],
    ])
    result = parse_excel(f)
    cells = result.sheets[0].rows[0].cells
    assert len(cells) == 2
    assert all(c.dates == [] for c in cells)


def test_parse_empty_group_before_date():
    """Колонка «Дата» без предшествующих типов — не создаёт пустую группу."""
    f = _make_xlsx([
        ["Наряд", "Дата", "ШЛ", "Дата"],
        ["", "", "Ящики", ""],
        ["Тест001", "2026-01-01", 5, "2026-10-08"],
    ])
    result = parse_excel(f)
    cells = result.sheets[0].rows[0].cells
    assert len(cells) == 1
    assert cells[0].research_type_code == "ШЛ"
    assert cells[0].dates == [date(2026, 10, 8)]


# ============================================================
# Шапка + тип тары (receiving-1a-2)
# ============================================================
def test_parse_header_row_after_two_title_rows():
    """Реальный файл: 2 строки шапки, затем заголовки, затем тип тары."""
    f = _make_xlsx([
        ["АКТ сдачи-приема", "", "", ""],
        ["Проб отработанных в Лаб1 отправляемых в Склад", "", "", ""],
        ["№ Наряд-заказа", "ПА", "ГД", "Дата"],
        ["", "Ящики", "Ящики", ""],
        ["TST00469", 3, 4, "2026-09-08"],
    ])
    result = parse_excel(f)
    assert len(result.errors) == 0
    assert len(result.sheets) == 1
    sheet = result.sheets[0]
    assert sheet.header_row_number == 3
    assert len(sheet.rows) == 1
    cells = sheet.rows[0].cells
    assert len(cells) == 2
    assert cells[0].container_type_name == "Ящики"
    assert cells[1].container_type_name == "Ящики"


def test_parse_container_type_per_column():
    """Тип тары разный для разных колонок (Ящики vs Биг-бэг)."""
    f = _make_xlsx([
        ["№ Наряд-заказа", "АД", "ГД", "Дата", "Хвосты дробления", "Дата"],
        ["", "Ящики", "Ящики", "", "Биг-бэг", ""],
        ["TST00470", 1, 3, "2026-09-11", 5, "2026-09-12"],
    ])
    result = parse_excel(f)
    cells = result.sheets[0].rows[0].cells
    ad = next(c for c in cells if c.research_type_code == "АД")
    gd = next(c for c in cells if c.research_type_code == "ГД")
    hd = next(c for c in cells if c.research_type_code == "Хвосты дробления")
    assert ad.container_type_name == "Ящики"
    assert gd.container_type_name == "Ящики"
    assert hd.container_type_name == "Биг-бэг"


def test_parse_missing_container_type_is_error():
    """Колонка типа без типа тары → ошибка MISSING_CONTAINER_TYPE."""
    f = _make_xlsx([
        ["№ Наряд-заказа", "ШЛ", "ХА", "Дата"],
        ["", "Ящики", "", ""],  # для ХА тип тары не указан
        ["Тест001", 5, 3, "2026-10-08"],
    ])
    result = parse_excel(f)
    missing_errors = [
        e for e in result.errors
        if e.get("code") == "MISSING_CONTAINER_TYPE"
    ]
    assert len(missing_errors) == 1
    # ХА отброшен, ШЛ остался.
    cells = result.sheets[0].rows[0].cells
    assert len(cells) == 1
    assert cells[0].research_type_code == "ШЛ"


def test_parse_no_container_type_row_at_all():
    """Строки типа тары нет вообще — все колонки с ошибкой."""
    f = _make_xlsx([
        ["№ Наряд-заказа", "ШЛ", "ХА", "Дата"],
        ["Тест001", 5, 3, "2026-10-08"],
    ])
    result = parse_excel(f)
    missing_errors = [
        e for e in result.errors
        if e.get("code") == "MISSING_CONTAINER_TYPE"
    ]
    assert len(missing_errors) == 2


def test_parse_container_type_row_shorter_than_headers():
    """Строка типа тары короче строки заголовков — лишние колонки с ошибкой."""
    f = _make_xlsx([
        ["№ Наряд-заказа", "ШЛ", "ХА", "ИЗ", "Дата"],
        ["", "Ящики", "Ящики"],  # ИЗ без типа тары
        ["Тест001", 5, 3, 1, "2026-10-08"],
    ])
    result = parse_excel(f)
    missing_errors = [
        e for e in result.errors
        if e.get("code") == "MISSING_CONTAINER_TYPE"
    ]
    assert len(missing_errors) == 1
    cells = result.sheets[0].rows[0].cells
    assert len(cells) == 2